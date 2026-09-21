"""네이버 검색 결과(반도체)에서 신문기사 링크를 수집하고 기사 본문을 크롤링한다.

사용법:
    python news_crawler.py                 # 기본 URL, 상위 5건 출력 + news_result.json 저장
    python news_crawler.py -n 10           # 10건
    python news_crawler.py -q 인공지능      # 검색어 변경
    python news_crawler.py -u "<검색 URL>"  # URL 직접 지정

필요 패키지: pip install requests beautifulsoup4
"""
import argparse
import json
import re
import sys
import time
from urllib.parse import quote, urljoin, urlparse

import requests
from bs4 import BeautifulSoup

DEFAULT_URL = (
    "https://search.naver.com/search.naver?where=nexearch&sm=top_hty&fbm=0"
    "&ie=utf8&query=%EB%B0%98%EB%8F%84%EC%B2%B4&ackey=uvpy7u54"
)
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9",
}
TIMEOUT = 10

# 언론사마다 본문 컨테이너가 다르므로 자주 쓰이는 선택자를 순서대로 시도한다
BODY_SELECTORS = [
    "#dic_area",                    # 네이버 뉴스
    "#articleBodyContents",
    "#newsct_article",
    "[itemprop='articleBody']",
    "#article-view-content-div",
    "#articleBody",
    "#article_body",
    ".article_body",
    ".article-body",
    ".news_body",
    "#newsContent",
    "article",
]
NOISE_TAGS = ["script", "style", "figure", "figcaption", "aside", "iframe",
              "noscript", "form", "button", "nav", "footer", "header"]


def get_html(url):
    resp = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
    resp.raise_for_status()
    # 일부 언론사는 charset 헤더가 없어 EUC-KR 등이 깨지므로 추정값을 사용한다
    if resp.encoding is None or resp.encoding.lower() in ("iso-8859-1", "ascii"):
        resp.encoding = resp.apparent_encoding
    return resp.text


def is_article_link(url):
    """네이버 자체 서비스(블로그, 클립 등)를 제외한 기사 링크인지 판단."""
    p = urlparse(url)
    if p.scheme not in ("http", "https"):
        return False
    host = p.netloc.lower()
    if host.endswith("naver.com") or host.endswith("naver.net"):
        return host in ("news.naver.com", "n.news.naver.com") and "/article/" in p.path
    return len(p.path.strip("/")) > 0


def gdid_of(a):
    """data-nlog-params 의 gdid: 같은 기사의 제목 링크와 네이버뉴스 링크를 이어주는 키."""
    try:
        return json.loads(a.get("data-nlog-params", "{}")).get("gdid", "")
    except ValueError:
        return ""


def search_articles_by_area(soup, search_url):
    """뉴스 목록의 data-nlog-area 속성(nws_all.*.tit / nws_all.*.nav)으로 기사를 찾는다."""
    nav = {}  # gdid -> 네이버뉴스 URL (원문 본문을 못 읽을 때의 대체 경로)
    for a in soup.select('a[data-nlog-area^="nws_all."][data-nlog-area$=".nav"]'):
        nav.setdefault(gdid_of(a), urljoin(search_url, a["href"]))

    results, seen = [], set()
    for a in soup.select('a[data-nlog-area^="nws_all."][data-nlog-area$=".tit"]'):
        href = urljoin(search_url, a["href"])
        if href in seen:
            continue
        seen.add(href)
        span = a.select_one("span.sds-comps-text")  # "새 창 열림" 안내 span 제외
        title = (span or a).get_text(" ", strip=True)
        results.append({"title": title, "url": href, "naver_url": nav.get(gdid_of(a), "")})
    return results


def search_articles(search_url):
    """검색 결과 페이지에서 (제목, 기사 URL) 목록을 반환한다."""
    soup = BeautifulSoup(get_html(search_url), "html.parser")
    results = search_articles_by_area(soup, search_url)
    if results:
        return results
    # 네이버 마크업이 바뀐 경우를 위한 대체 방법: 링크 주소로 기사를 추정
    seen, results = set(), []
    for a in soup.select("a[href]"):
        href = urljoin(search_url, a["href"])
        title = re.sub(r"\s*새 창 열림$", "", a.get_text(" ", strip=True))
        if title.startswith("네이버뉴스"):  # 언론사 링크 옆의 "네이버뉴스" 버튼: 제목은 본문 수집 후 채움
            title = ""
        # 같은 기사 링크가 제목/요약으로 두 번 나오므로 처음(제목)만 사용
        if href in seen or not is_article_link(href):
            continue
        if title and len(title) < 12:
            continue
        if not title and "news.naver.com" not in href:
            continue
        seen.add(href)
        results.append({"title": title, "url": href})
    return results


def clean_text(text):
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_body(soup):
    for sel in BODY_SELECTORS:
        node = soup.select_one(sel)
        if node:
            for t in node.select(",".join(NOISE_TAGS)):
                t.decompose()
            text = clean_text(node.get_text("\n", strip=True))
            if len(text) >= 100:
                return text
    # 마지막 수단: <p> 문단이 가장 많이 모인 부모 요소를 본문으로 본다
    best, best_len = None, 0
    for parent in {p.parent for p in soup.find_all("p") if p.parent}:
        length = sum(len(p.get_text(strip=True)) for p in parent.find_all("p", recursive=False))
        if length > best_len:
            best, best_len = parent, length
    if best is not None:
        return clean_text("\n".join(p.get_text(" ", strip=True)
                                    for p in best.find_all("p", recursive=False)))
    return ""


def meta(soup, *names):
    for n in names:
        tag = soup.find("meta", attrs={"property": n}) or soup.find("meta", attrs={"name": n})
        if tag and tag.get("content"):
            return tag["content"].strip()
    return ""


def crawl_article(url):
    soup = BeautifulSoup(get_html(url), "html.parser")
    title = meta(soup, "og:title", "twitter:title") or (
        soup.title.get_text(strip=True) if soup.title else "")
    published = meta(soup, "article:published_time", "og:regDate", "pubdate")
    press = meta(soup, "og:site_name", "article:author")
    body = extract_body(soup)
    summary = meta(soup, "og:description", "description")
    if len(body) < 100 and len(summary) > len(body):  # 본문이 막힌 경우 요약이라도 저장
        body = summary
    return {"title": title, "press": press, "published": published,
            "url": url, "content": body}


EXCEL_CELL_LIMIT = 32000  # 엑셀 한 셀의 최대 글자 수(32,767)보다 조금 작게
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def save_excel(articles, path):
    """기사 목록을 엑셀(.xlsx)로 저장한다. (pip install openpyxl)"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def clean(v):  # 엑셀이 허용하지 않는 제어문자 제거 + 셀 길이 제한
        return _ILLEGAL_XML.sub("", str(v or ""))[:EXCEL_CELL_LIMIT]

    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"
    headers = ["번호", "제목", "언론사", "발행일", "URL", "본문"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3B5BDB")
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, a in enumerate(articles, 1):
        published = clean(a.get("published")).replace("T", " ")[:19]
        ws.append([i, clean(a.get("title")), clean(a.get("press")), published,
                   clean(a.get("url")), clean(a.get("content"))])
        link = ws.cell(row=i + 1, column=5)
        if link.value.startswith(("http://", "https://")) and len(link.value) <= 2000:
            link.hyperlink = link.value
            link.font = Font(color="0563C1", underline="single")
        for col in (2, 6):
            ws.cell(row=i + 1, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (1, 3, 4, 5):
            ws.cell(row=i + 1, column=col).alignment = Alignment(vertical="top")

    for idx, width in enumerate([6, 45, 14, 20, 40, 100], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    # Windows 콘솔에서 한글이 깨지지 않도록
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    ap = argparse.ArgumentParser(description="네이버 검색 결과 뉴스 크롤러")
    ap.add_argument("-u", "--url", default=DEFAULT_URL, help="검색 결과 URL")
    ap.add_argument("-q", "--query", help="검색어 (지정하면 --url보다 우선)")
    ap.add_argument("-n", "--num", type=int, default=5, help="수집할 기사 수")
    ap.add_argument("-o", "--out", default="news_result.json", help="저장할 JSON 파일")
    ap.add_argument("-x", "--xlsx", nargs="?", const="news_result.xlsx", metavar="파일",
                    help="엑셀(.xlsx)로도 저장 (파일명 생략 시 news_result.xlsx)")
    args = ap.parse_args()

    url = (f"https://search.naver.com/search.naver?where=nexearch&ie=utf8&query={quote(args.query)}"
           if args.query else args.url)

    print(f"검색 페이지 요청: {url}\n")
    links = search_articles(url)
    if not links:
        print("기사 링크를 찾지 못했습니다. 네이버 페이지 구조가 바뀌었을 수 있습니다.")
        return

    articles, seen_titles = [], set()
    for i, item in enumerate(links, 1):
        if len(articles) >= args.num:
            break
        print(f"[{i}/{len(links)}] {item['title'] or item['url']}")
        try:
            art = crawl_article(item["url"])
            # 원문 사이트가 본문을 막은 경우 같은 기사의 네이버뉴스 사본으로 대체
            if len(art["content"]) < 150 and item.get("naver_url"):
                alt = crawl_article(item["naver_url"])
                if len(alt["content"]) > len(art["content"]):
                    art["content"] = alt["content"]
                    art["content_source"] = item["naver_url"]
        except requests.RequestException as e:
            print(f"    실패: {e}\n")
            continue
        if art["title"] in seen_titles:  # 언론사 원문과 네이버뉴스 사본이 함께 나오는 경우
            print("    (중복 기사, 건너뜀)\n")
            continue
        seen_titles.add(art["title"])
        print(f"    제목: {art['title']}")
        articles.append(art)
        preview = art["content"][:200].replace("\n", " ")
        print(f"    URL : {art['url']}")
        print(f"    언론사/일시: {art['press'] or '-'} / {art['published'] or '-'}")
        print(f"    본문({len(art['content'])}자): {preview}...\n")
        time.sleep(1)  # 서버에 부담을 주지 않도록 요청 간격을 둔다

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(articles, f, ensure_ascii=False, indent=2)
    print(f"{len(articles)}건을 {args.out}에 저장했습니다.")
    if args.xlsx:
        save_excel(articles, args.xlsx)
        print(f"{len(articles)}건을 {args.xlsx}에 저장했습니다.")


if __name__ == "__main__":
    main()
